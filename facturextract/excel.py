"""Generador del Excel mensual: solo lectura para el cliente, regenerado desde la base.

Idempotente (mismo contenido cada vez) y atómico (archivo temporal + reemplazo). Si el archivo
está abierto o bloqueado, el mes queda marcado como desactualizado y se reintenta después.
"""
from __future__ import annotations

import logging
import os
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import Reloj
from .db import Repo
from .modelo import TIPOS_DOCUMENTO, parse_fecha

log = logging.getLogger(__name__)

COLUMNAS = [
    "RUT Emisor", "Razón Social Emisor", "RUT Receptor", "Razón Social Receptor", "Tipo de Documento", "Folio",
    "Fecha Emisión", "Fecha Vencimiento", "Forma de Pago (documento)", "Cantidad", "Detalle Producto", "Precio Unitario",
    "Descuento", "Total Línea", "Monto Neto", "Monto Exento", "IVA", "Otros Impuestos", "Total Factura",
    "Estado Validación", "Observaciones", "Estado de Pago", "Fecha de Pago", "Tipo de Pago", "N° de Documento",
]
FORMATO_FECHA = "DD-MM-YYYY"
COLUMNAS_FECHA = {"Fecha Emisión", "Fecha Vencimiento", "Fecha de Pago"}


def _fecha(valor: str | None) -> date | None:
    return parse_fecha(valor) if valor else None


def fila_excel(f: dict, ln: dict | None) -> list:
    tipo_pago = (f["metodo_pago"] or "").capitalize()
    return [
        f["rut_emisor"], f["razon_social_emisor"], f["rut_receptor"], f["razon_social_receptor"],
        TIPOS_DOCUMENTO.get(f["tipo_documento"], f["tipo_documento"]), int(f["folio"]) if f["folio"] else None,
        _fecha(f["fecha_emision"]), _fecha(f["fecha_vencimiento"]), f["medio_pago"],
        ln["cantidad"] if ln else None, ln["detalle"] if ln else None, ln["precio_unitario"] if ln else None,
        ln["descuento"] if ln else None, ln["total_linea"] if ln else None,
        f["monto_neto"], f["monto_exento"], f["iva"], f["otros_impuestos"], f["total"],
        f["estado_validacion"], f["observaciones"], (f["estado_pago"] or "").capitalize(), _fecha(f["fecha_pago"]),
        tipo_pago,
        f["nro_documento"],
    ]


class ExcelMensual:
    def __init__(self, repo: Repo, carpeta: Path, reloj: Reloj | None = None):
        self.repo = repo
        self.carpeta = Path(carpeta)
        self.reloj = reloj or Reloj()

    def ruta_mes(self, mes: str) -> Path:
        anio, m = mes.split("-")
        return self.carpeta / f"facturas_{anio}_{m}.xlsx"

    def _construir(self, mes: str) -> Workbook:
        wb = Workbook()
        ws = wb.active
        ws.title = "Facturas"
        ws.append(COLUMNAS)
        for celda in ws[1]:
            celda.font = Font(bold=True)
        for fila in self.repo.filas_excel(mes):
            ws.append(fila_excel(fila["factura"], fila["linea"]))
        for i, nombre in enumerate(COLUMNAS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(nombre) + 4))
            if nombre in COLUMNAS_FECHA:
                for c in ws[get_column_letter(i)][1:]:
                    c.number_format = FORMATO_FECHA
        ws.freeze_panes = "A2"
        ws.protection.sheet = True  # solo lectura para el cliente (sin contraseña)
        return wb

    def regenerar(self, mes: str) -> str:
        """Devuelve 'ok' o 'desactualizado'. Nunca lanza por un archivo bloqueado."""
        tmp_ruta = None
        try:
            self.carpeta.mkdir(parents=True, exist_ok=True)
            wb = self._construir(mes)
            fd, tmp_nombre = tempfile.mkstemp(prefix=".tmp_", suffix=".xlsx", dir=self.carpeta)
            os.close(fd)
            tmp_ruta = Path(tmp_nombre)
            wb.save(tmp_ruta)
            os.replace(tmp_ruta, self.ruta_mes(mes))
            tmp_ruta = None
            self.repo.marcar_mes_actualizado(mes, self.reloj.ahora())
            return "ok"
        except OSError as e:  # archivo abierto en Excel, disco lleno, etc.
            log.warning("No se pudo escribir el Excel de %s: %s", mes, e)
            self.repo.marcar_mes_desactualizado(mes)
            return "desactualizado"
        finally:
            if tmp_ruta is not None:
                try:
                    tmp_ruta.unlink()
                except OSError:
                    pass

    def reintentar_pendientes(self) -> dict[str, str]:
        return {mes: self.regenerar(mes) for mes in self.repo.meses_desactualizados()}
