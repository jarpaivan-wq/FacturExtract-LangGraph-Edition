import os
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from facturextract import excel as excel_mod
from facturextract.db import DecisionInvalida, Duplicada
from facturextract.excel import COLUMNAS, ExcelMensual
from tests.factura_ejemplo import factura

HOY = date(2026, 10, 12)
AHORA = datetime(2026, 10, 12, 9, 0)


def guardar(repo, fid="a", hoy=HOY, manual=False, **cambios):
    repo.crear_factura(fid, ["x.jpg"], AHORA)
    return repo.guardar_factura(fid, factura(**cambios), [], hoy, manual)


def seg(repo):
    return repo._leer("SELECT * FROM seguimientos")


def test_guardar_crea_factura_lineas_y_seguimiento(repo):
    res = guardar(repo, fecha_vencimiento="2026-11-04")
    assert res == {"mes": "2026-10", "seguimiento": True}
    f = repo.obtener("a")
    assert f["estado"] == "guardada" and f["estado_pago"] == "pendiente" and f["estado_validacion"] == "Válida"
    assert len(repo.lineas_de("a")) == 2
    s = seg(repo)
    assert len(s) == 1 and s[0]["proxima_alerta"] == "2026-11-01"   # vencimiento - 3 días corridos


def test_vencimiento_cercano_o_vencida_alerta_inmediata(repo):
    guardar(repo, "a", fecha_vencimiento="2026-10-14", folio="1")
    guardar(repo, "b", fecha_vencimiento="2026-10-01", folio="2")
    assert {s["proxima_alerta"] for s in seg(repo)} == {"2026-10-12"}


def test_contado_y_nota_no_generan_seguimiento_pero_una_pendiente_de_mes_anterior_si(repo):
    assert not guardar(repo, "a", folio="1", pago_sugerido="contado", fecha_vencimiento="2026-11-04")["seguimiento"]
    f = repo.obtener("a")
    assert (f["estado_pago"], f["metodo_pago"], f["fecha_pago"]) == ("pagada", "contado", "2026-10-12")
    assert not guardar(repo, "b", folio="2", tipo_documento="nota_credito")["seguimiento"]
    assert repo.obtener("b")["estado_pago"] == "no aplica"
    assert guardar(repo, "c", folio="3", fecha_emision="2026-09-20")["seguimiento"]      # regla nueva: alerta igual
    assert len(seg(repo)) == 1 and repo.obtener("c")["mes_emision"] == "2026-09"          # y va al Excel de septiembre


def test_duplicado_bloqueado_por_restriccion_unica(repo):
    guardar(repo, "a")
    assert repo.existe_guardada(factura()["rut_emisor"], "1001", "factura")
    with pytest.raises(Duplicada):
        guardar(repo, "b")
    assert repo.obtener("b")["estado"] == "registrada"       # la transacción se revirtió
    assert len(repo._leer("SELECT * FROM lineas WHERE factura_id='b'")) == 0
    assert len(seg(repo)) == 1                                  # sin alerta duplicada
    # mismo folio pero otro tipo de documento: no es duplicado
    guardar(repo, "c", tipo_documento="nota_credito", fecha_vencimiento=None)


def test_transaccion_atomica_si_falla_a_medio_camino(repo, monkeypatch):
    repo.crear_factura("a", [], AHORA)
    malo = factura()
    malo["lineas"][1]["detalle"] = object()     # no se puede insertar -> falla tras actualizar la cabecera
    with pytest.raises(Exception):
        repo.guardar_factura("a", malo, [], HOY)
    assert repo.obtener("a")["estado"] == "registrada" and repo.lineas_de("a") == []


def test_pago_con_y_sin_cheque(repo):
    guardar(repo, "a", folio="1")
    guardar(repo, "b", folio="2")
    with pytest.raises(DecisionInvalida):
        repo.registrar_pago("a", HOY, "cheque", "  ")                  # solo el cheque exige N° de documento
    assert repo.registrar_pago("a", HOY, "cheque", "5555") == "2026-10"
    f = repo.obtener("a")
    assert (f["estado_pago"], f["metodo_pago"], f["fecha_pago"], f["nro_documento"]) == (
        "pagada", "cheque", "2026-10-12", "5555")
    repo.registrar_pago("b", HOY, "transferencia", "999")             # el N° se ignora si no es cheque
    f = repo.obtener("b")
    assert (f["metodo_pago"], f["nro_documento"]) == ("transferencia", None)
    with pytest.raises(DecisionInvalida):
        repo.registrar_pago("b", HOY, "efectivo", None)               # métodos válidos: contado, transferencia, cheque
    assert {s["estado"] for s in seg(repo)} == {"pagada"}


def test_pagar_una_factura_ya_pagada_se_rechaza(repo):
    guardar(repo, "a", pago_sugerido="contado")
    with pytest.raises(DecisionInvalida):
        repo.registrar_pago("a", HOY, "contado", None)


def test_estado_validacion_con_advertencias_y_manual(repo):
    repo.crear_factura("a", [], AHORA)
    adv = [{"codigo": "V09", "mensaje": "Medio raro", "campo": None, "gravedad": "advertencia"}]
    repo.guardar_factura("a", factura(), adv, HOY)
    assert repo.obtener("a")["estado_validacion"].startswith("Válida con advertencias")
    guardar(repo, "b", folio="7", manual=True)
    assert repo.obtener("b")["estado_validacion"] == "Aprobada manualmente"


def leer_excel(ruta):
    ws = load_workbook(ruta).active
    return [[c.value for c in fila] for fila in ws.iter_rows()]


def test_excel_columnas_filas_por_linea_y_pago(repo, r):
    guardar(repo, "a")
    repo.registrar_pago("a", HOY, "cheque", "777")
    ex = ExcelMensual(repo, r.excel)
    assert ex.regenerar("2026-10") == "ok"
    filas = leer_excel(ex.ruta_mes("2026-10"))
    assert filas[0] == COLUMNAS
    assert len(filas) == 3                                       # 2 líneas
    i = COLUMNAS.index
    for fila in filas[1:]:                                        # el pago va en TODAS las filas
        assert (fila[i("Estado de Pago")], fila[i("Tipo de Pago")], fila[i("N° de Documento")]) == ("Pagada", "Cheque", "777")
    assert repo.meses_desactualizados() == []


def test_excel_idempotente_y_sin_temporales(repo, r):
    guardar(repo, "a")
    ex = ExcelMensual(repo, r.excel)
    ex.regenerar("2026-10")
    a = leer_excel(ex.ruta_mes("2026-10"))
    ex.regenerar("2026-10")
    assert leer_excel(ex.ruta_mes("2026-10")) == a
    assert [p.name for p in r.excel.iterdir()] == ["facturas_2026_10.xlsx"]


def test_excel_por_mes_de_emision(repo, r):
    guardar(repo, "a", folio="1", fecha_emision="2026-09-20", fecha_vencimiento="2026-10-20")
    guardar(repo, "b", folio="2")
    ex = ExcelMensual(repo, r.excel)
    ex.regenerar("2026-09"); ex.regenerar("2026-10")
    assert len(leer_excel(ex.ruta_mes("2026-09"))) == 3 and len(leer_excel(ex.ruta_mes("2026-10"))) == 3


def test_excel_bloqueado_marca_desactualizado_y_reintenta(repo, r, monkeypatch):
    guardar(repo, "a")
    ex = ExcelMensual(repo, r.excel)
    real = os.replace

    def bloqueado(*a, **k):
        raise PermissionError("archivo en uso")

    monkeypatch.setattr(excel_mod.os, "replace", bloqueado)
    assert ex.regenerar("2026-10") == "desactualizado"
    assert repo.meses_desactualizados() == ["2026-10"]
    assert not ex.ruta_mes("2026-10").exists()
    assert list(r.excel.iterdir()) == []                          # sin temporales huérfanos
    monkeypatch.setattr(excel_mod.os, "replace", real)
    assert ex.reintentar_pendientes() == {"2026-10": "ok"}
    assert repo.meses_desactualizados() == [] and ex.ruta_mes("2026-10").exists()


def test_excel_bloqueado_conserva_version_anterior(repo, r, monkeypatch):
    guardar(repo, "a")
    ex = ExcelMensual(repo, r.excel)
    ex.regenerar("2026-10")
    antes = leer_excel(ex.ruta_mes("2026-10"))
    guardar(repo, "b", folio="2")
    monkeypatch.setattr(excel_mod.os, "replace", lambda *a, **k: (_ for _ in ()).throw(PermissionError()))
    assert ex.regenerar("2026-10") == "desactualizado"
    assert leer_excel(ex.ruta_mes("2026-10")) == antes


def test_pago_no_cambia_el_vencimiento_ni_la_alerta_original(repo):
    guardar(repo, "a", fecha_vencimiento="2026-11-04")
    repo.registrar_pago("a", HOY, "transferencia", None)
    f = repo.obtener("a")
    assert f["fecha_vencimiento"] == "2026-11-04"                      # el vencimiento original se conserva
    assert (f["estado_pago"], f["metodo_pago"]) == ("pagada", "transferencia")


def test_campos_de_pago_estan_separados_en_el_esquema(repo):
    cols = {r["name"] for r in repo._leer("PRAGMA table_info(facturas)")}
    assert {"estado_pago", "metodo_pago", "nro_documento", "fecha_pago", "fecha_vencimiento"} <= cols
    assert "con_cheque" not in cols
    assert "proxima_alerta" in {r["name"] for r in repo._leer("PRAGMA table_info(seguimientos)")}


def test_guardar_con_configuracion_explicita_de_pago(repo):
    repo.crear_factura("a", [], AHORA)
    pago = {"estado_pago": "pendiente", "metodo_pago": None, "nro_documento": None, "fecha_alerta": "2026-10-20"}
    repo.guardar_factura("a", factura(fecha_vencimiento="2026-11-04"), [], HOY, False, pago)
    assert seg(repo)[0]["proxima_alerta"] == "2026-10-20"
    assert repo.obtener("a")["fecha_vencimiento"] == "2026-11-04"


def test_migracion_desde_esquema_v1(tmp_path):
    import sqlite3

    from facturextract.db import Repo
    ruta = tmp_path / "viejo.db"
    con = sqlite3.connect(ruta)
    con.executescript("""
        CREATE TABLE facturas (id TEXT PRIMARY KEY, estado TEXT NOT NULL, creada_en TEXT NOT NULL,
            imagenes TEXT NOT NULL DEFAULT '[]', extraido TEXT, hallazgos TEXT, mensaje TEXT,
            tipo_documento TEXT, rut_emisor TEXT, razon_social_emisor TEXT, rut_receptor TEXT,
            razon_social_receptor TEXT, folio TEXT, fecha_emision TEXT, fecha_vencimiento TEXT, medio_pago TEXT,
            monto_neto INTEGER, monto_exento INTEGER, iva INTEGER, otros_impuestos INTEGER, total INTEGER,
            observaciones TEXT, aprobada_manual INTEGER NOT NULL DEFAULT 0, estado_validacion TEXT,
            mes_emision TEXT, estado_pago TEXT, fecha_pago TEXT, con_cheque INTEGER, nro_documento TEXT);
        INSERT INTO facturas (id, estado, creada_en, estado_pago, con_cheque, fecha_pago, nro_documento)
            VALUES ('c','guardada','2026-10-01T10:00:00','pagada',1,'2026-10-02','77'),
                   ('t','guardada','2026-10-01T10:00:00','pagada',0,'2026-10-02',NULL),
                   ('k','guardada','2026-10-01T10:00:00','contado',NULL,NULL,NULL),
                   ('p','guardada','2026-10-01T10:00:00','pendiente',NULL,NULL,NULL);
    """)
    con.commit()
    con.close()
    r = Repo(ruta)
    g = {x["id"]: x for x in (r.obtener(i) for i in "ctkp")}
    assert (g["c"]["estado_pago"], g["c"]["metodo_pago"], g["c"]["nro_documento"]) == ("pagada", "cheque", "77")
    assert (g["t"]["estado_pago"], g["t"]["metodo_pago"]) == ("pagada", "transferencia")
    assert (g["k"]["estado_pago"], g["k"]["metodo_pago"], g["k"]["fecha_pago"]) == ("pagada", "contado", "2026-10-01")
    assert (g["p"]["estado_pago"], g["p"]["metodo_pago"]) == ("pendiente", None)
    assert "con_cheque" not in g["c"]
    Repo(ruta)                                                         # reabrir es idempotente
