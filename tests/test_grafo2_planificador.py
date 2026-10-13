from datetime import date, datetime

import pytest

from facturextract import grafo2
from facturextract.config import RelojFijo
from facturextract.excel import ExcelMensual
from facturextract.extractor import ExtractorFalso
from facturextract.grafo1 import Dependencias, abrir_checkpointer
from facturextract.planificador import Planificador
from tests.factura_ejemplo import factura

AHORA = datetime(2026, 10, 12, 9, 30)


class Entorno:
    def __init__(self, repo, r, reloj):
        self.repo, self.r, self.reloj = repo, r, reloj
        self.cp = abrir_checkpointer(r.checkpoints_db)
        self.deps = Dependencias(repo, ExtractorFalso([]), ExcelMensual(repo, r.excel, reloj), reloj, r)
        self.g2 = grafo2.construir_grafo2(self.deps, self.cp)
        self.plan = Planificador(repo, self.g2)


@pytest.fixture
def env(repo, r, reloj):
    return Entorno(repo, r, reloj)


def registrar(repo, fid="a", hoy=date(2026, 10, 12), **cambios):
    repo.crear_factura(fid, [], AHORA)
    repo.guardar_factura(fid, factura(**cambios), [], hoy)
    return repo._leer("SELECT id FROM seguimientos WHERE factura_id=?", (fid,))[0]["id"]


def alerta(repo, sid):
    return repo.obtener_seguimiento(sid)


def test_no_alerta_antes_de_3_dias_y_si_ese_dia(env):
    sid = registrar(env.repo, fecha_emision="2026-10-05", fecha_vencimiento="2026-10-20")   # alerta 17-oct
    assert env.plan.revisar(datetime(2026, 10, 16, 12, 0))["lanzados"] == 0
    r = env.plan.revisar(datetime(2026, 10, 17, 9, 0))
    assert r["lanzados"] == 1 and alerta(env.repo, sid)["estado"] == "esperando_decision"


def test_vencida_al_registrar_alerta_inmediata(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    r = env.plan.revisar(AHORA)
    assert r["lanzados"] == 1 and r["abrir_ventana"] and alerta(env.repo, sid)["estado"] == "esperando_decision"


def test_antes_de_la_hora_configurada_no_hay_alerta(env):
    registrar(env.repo, fecha_vencimiento="2026-10-01")
    assert env.plan.revisar(datetime(2026, 10, 12, 8, 59))["lanzados"] == 0
    env.repo.set_ajuste("hora_alerta", "07:00")
    assert env.plan.revisar(datetime(2026, 10, 12, 8, 59))["lanzados"] == 1


def test_un_run_por_seguimiento_y_dia_y_ventana_una_vez(env):
    registrar(env.repo, fecha_vencimiento="2026-10-01")
    r1 = env.plan.revisar(AHORA)
    r2 = env.plan.revisar(datetime(2026, 10, 12, 9, 45))
    assert (r1["lanzados"], r1["abrir_ventana"]) == (1, True)
    assert (r2["lanzados"], r2["abrir_ventana"], r2["esperando"]) == (0, False, 1)
    assert len(env.repo.seguimientos_esperando()) == 1


def test_cerrar_reaparece_al_dia_siguiente(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    out = env.plan.decidir(sid, {"accion": "cerrar"})
    assert out["resultado"] == "cerrada"
    s = alerta(env.repo, sid)
    assert (s["estado"], s["proxima_alerta"]) == ("pendiente", "2026-10-13")
    assert env.plan.revisar(datetime(2026, 10, 12, 15, 0))["lanzados"] == 0     # no vuelve hoy
    r = env.plan.revisar(datetime(2026, 10, 13, 9, 0))
    assert r["lanzados"] == 1 and r["abrir_ventana"]                             # reaparece mañana


def test_posponer_a_fecha_elegida(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    out = env.plan.decidir(sid, {"accion": "posponer", "fecha": "20-10-2026"})
    assert out["resultado"] == "pospuesta" and alerta(env.repo, sid)["proxima_alerta"] == "2026-10-20"
    assert env.plan.revisar(datetime(2026, 10, 19, 10, 0))["lanzados"] == 0
    assert env.plan.revisar(datetime(2026, 10, 20, 10, 0))["lanzados"] == 1


def test_posponer_exige_fecha_posterior_a_hoy_y_repregunta(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    for dec in ({"accion": "posponer", "fecha": "12-10-2026"}, {"accion": "posponer"},
                {"accion": "posponer", "fecha": "ayer"}, {"accion": "volar"}):
        out = env.plan.decidir(sid, dec)
        assert out["__interrupt__"][0].value["error"], dec
        assert alerta(env.repo, sid)["estado"] == "esperando_decision"
    assert env.plan.decidir(sid, {"accion": "posponer", "fecha": "2026-10-13"})["resultado"] == "pospuesta"


def test_pagada_sin_cheque(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    out = env.plan.decidir(sid, {"accion": "pagada", "metodo": "transferencia"})
    assert out["resultado"] == "pagada" and out["excel"] == "ok"
    f = env.repo.obtener("a")
    assert (f["estado_pago"], f["metodo_pago"], f["fecha_pago"]) == ("pagada", "transferencia", "2026-10-12")
    assert f["fecha_vencimiento"] == "2026-10-01"                   # pagar no cambia el vencimiento
    assert alerta(env.repo, sid)["estado"] == "pagada"
    assert env.plan.revisar(datetime(2026, 10, 30, 9, 0))["lanzados"] == 0


def test_pagada_con_cheque_exige_numero(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    out = env.plan.decidir(sid, {"accion": "pagada", "metodo": "cheque", "nro_documento": " "})
    assert "N° de documento" in out["__interrupt__"][0].value["error"]
    assert env.repo.obtener("a")["estado_pago"] == "pendiente"
    out = env.plan.decidir(sid, {"accion": "pagada", "metodo": "cheque", "nro_documento": "8821"})
    assert out["resultado"] == "pagada"
    from openpyxl import load_workbook
    ws = load_workbook(env.deps.excel.ruta_mes("2026-10")).active
    cab = [c.value for c in ws[1]]
    for fila in list(ws.iter_rows(min_row=2, values_only=True)):
        assert fila[cab.index("N° de Documento")] == "8821" and fila[cab.index("Tipo de Pago")] == "Cheque"


def test_contado_y_notas_no_alertan(env):
    env.repo.crear_factura("c", [], AHORA)
    env.repo.guardar_factura("c", factura(pago_sugerido="contado", fecha_vencimiento="2026-10-01"), [], date(2026, 10, 12))
    assert env.plan.revisar(AHORA) == {"lanzados": 0, "esperando": 0, "abrir_ventana": False}


def test_decidir_sin_pausa_no_hace_nada(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-20")
    assert env.plan.decidir(sid, {"accion": "cerrar"})["resultado"] == "no_esperando"


def test_pagada_desde_la_vista_mientras_espera_la_alerta(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    env.repo.registrar_pago("a", date(2026, 10, 12), "contado", None)      # desde la vista de facturas
    assert env.repo.seguimientos_esperando() == []
    assert env.plan.decidir(sid, {"accion": "cerrar"})["resultado"] == "no_esperando"


def test_reanudar_pausa_tras_reiniciar_proceso(repo, r, reloj):
    e1 = Entorno(repo, r, reloj)
    sid = registrar(repo, fecha_vencimiento="2026-10-01")
    e1.plan.revisar(AHORA)
    e1.cp.conn.close()
    e2 = Entorno(repo, r, reloj)                       # proceso nuevo
    assert e2.plan.retomar_al_iniciar(date(2026, 10, 12)) == 0
    assert e2.plan.decidir(sid, {"accion": "cerrar"})["resultado"] == "cerrada"


def test_checkpoints_perdidos_se_recrean_desde_la_base(repo, r, reloj):
    e1 = Entorno(repo, r, reloj)
    sid = registrar(repo, fecha_vencimiento="2026-10-01")
    e1.plan.revisar(AHORA)
    e1.cp.conn.close()
    r.checkpoints_db.unlink()                          # checkpoints.db desechable
    for ext in ("-wal", "-shm"):
        p = r.checkpoints_db.with_name(r.checkpoints_db.name + ext)
        if p.exists():
            p.unlink()
    e2 = Entorno(repo, r, reloj)
    assert e2.plan.retomar_al_iniciar(date(2026, 10, 12)) == 1
    assert alerta(repo, sid)["estado"] == "pendiente"
    assert e2.plan.revisar(AHORA)["lanzados"] == 1     # se vuelve a lanzar y a mostrar


def test_pagada_exige_metodo_valido_y_solo_cheque_pide_numero(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    for dec in ({"accion": "pagada"}, {"accion": "pagada", "metodo": "efectivo"}):
        assert env.plan.decidir(sid, dec)["__interrupt__"][0].value["error"], dec
    out = env.plan.decidir(sid, {"accion": "pagada", "metodo": "contado"})      # contado: sin N° de documento
    assert out["resultado"] == "pagada" and env.repo.obtener("a")["nro_documento"] is None


def test_posponer_no_cambia_el_vencimiento_real(env):
    sid = registrar(env.repo, fecha_vencimiento="2026-10-01")
    env.plan.revisar(AHORA)
    env.plan.decidir(sid, {"accion": "posponer", "fecha": "2026-10-25"})
    assert env.repo.obtener("a")["fecha_vencimiento"] == "2026-10-01"
    assert alerta(env.repo, sid)["proxima_alerta"] == "2026-10-25"
